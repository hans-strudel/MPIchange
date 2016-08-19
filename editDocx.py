import sys
reload(sys)
sys.setdefaultencoding('utf-8')

import os
from os import listdir
from docx import Document
from docx.shared import Pt
from docx.package import Package
import gc

path = os.path.abspath(sys.argv[1])
files = listdir(path)
assy = sys.argv[2]
#in_file = os.path.abspath(sys.argv[1])
#out_file = os.path.abspath(sys.argv[2])
#assyNum = sys.argv[3]
#assyRev = sys.argv[4]

def remove_row(table, row):
	tbl = table._tbl
	tr = row._tr
	tbl.remove(tr)
	
def remove_cell(row, n):
	tr = row._tr
	tr.remove(tr[n]) 
	
for i in range(0, len(files)):
	in_file = path + '\\' + files[i]
	print files[i]
	out_file = os.path.abspath('mceMPI//MPI-' + files[i])
	document_part = Package.open(in_file)
	doc = document_part.main_document_part.document
	
	table = doc.tables[0]
	
	if ("Normal" in doc.styles): # set styling
		doc.styles['Normal'].font.name = 'Batang'
		doc.styles['Normal'].font.bold = True
		doc.styles['Normal'].font.size = Pt(12)
	else: # use deprecated backup method
		doc.styles[0].font.name = 'Batang'
		doc.styles[0].font.bold = True
		doc.styles[0].font.size = Pt(12)

	if (len(table.rows) > 3):
		remove_row(table, table.rows[3])	
	cols = len(table.columns)
	rows = len(table.rows)
	assyGuess = assy + table.cell(2,1).text.strip() + "-TK"
	
	
	assyNum = raw_input(assyGuess + " : ") or assyGuess
	table.cell(2,1).text =  assyNum
	descr = table.cell(1,1).text.strip()


	assyRevGuess = table.cell(1,3).text.strip() + "-A"
	if (table.rows[1].cells[1].text != table.rows[1].cells[2].text and table.rows[1].cells[1].text != table.rows[1].cells[2].text and table.rows[1].cells[2].text.find('FAB') < 0):
		
		## 4 cells in second col all descr
		table.cell(2,2).text = table.cell(1,2).text.strip()
		table.cell(1,2).text = ""
		assyRev = raw_input(assyRevGuess + " : ") or assyRevGuess
		table.cell(2,3).text =  assyRev
		table.cell(1,3).text = ""
		table.cell(1,2).merge(table.cell(1,3))
		table.cell(1,1).merge(table.cell(1,2))
	elif (table.rows[2].cells[2].text.find('FAB') > -1):
		print 'yes'
		temp = table.rows[1].cells[2].text.strip()
		assyRevGuess = table.cell(2,3).text.strip() + "-A"
		assyRev = raw_input(assyRevGuess + " : ") or assyRevGuess
		#table.rows[1].cells[2].text = table.cell(2,2).text.strip()
		#table.rows[1].cells[3].text = table.cell(2,3).text.strip()
		table.rows[1].cells[2].text = ""
		table.rows[1].cells[3].text = ""
		table.rows[2].cells[2].text = temp
		table.rows[2].cells[3].text = assyRev
	else:
		assyRevGuess = table.cell(2,3).text.strip() + "-A"
		assyRev = raw_input(assyRevGuess + " : ") or assyRevGuess
		table.cell(2,3).text = assyRev
		
	table.cell(1,1).text = descr
	g = gc.get_referrers(document_part)
	length = len(g)
	r = 0
	#for x in range(0, length-1):	
	#	print g[x]["_partname"]
	for a in range(0, len(g)-1):
		#print "Length of g ", len(g)
		#print a
		if ("_partname" in g[a]):
			if (g[a]["_partname"].find("header1.xml") > -1):
				g[a]["_blob"] = '<?xml version="1.0" encoding="UTF-8" standalone="yes"?><w:hdr xmlns:wpc="http://schemas.microsoft.com/office/word/2010/wordprocessingCanvas" xmlns:mc="http://schemas.openxmlformats.org/markup-compatibility/2006" xmlns:o="urn:schemas-microsoft-com:office:office" xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships" xmlns:m="http://schemas.openxmlformats.org/officeDocument/2006/math" xmlns:v="urn:schemas-microsoft-com:vml" xmlns:wp14="http://schemas.microsoft.com/office/word/2010/wordprocessingDrawing" xmlns:wp="http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing" xmlns:w10="urn:schemas-microsoft-com:office:word" xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" xmlns:w14="http://schemas.microsoft.com/office/word/2010/wordml" xmlns:w15="http://schemas.microsoft.com/office/word/2012/wordml" xmlns:wpg="http://schemas.microsoft.com/office/word/2010/wordprocessingGroup" xmlns:wpi="http://schemas.microsoft.com/office/word/2010/wordprocessingInk" xmlns:wne="http://schemas.microsoft.com/office/word/2006/wordml" xmlns:wps="http://schemas.microsoft.com/office/word/2010/wordprocessingShape" mc:Ignorable="w14 w15 wp14"><w:p w:rsidR="008A5523" w:rsidRDefault="008A5523" w:rsidP="00AF4909"><w:pPr><w:pStyle w:val="Header"/><w:jc w:val="both"/></w:pPr><w:proofErr w:type="spellStart"/><w:r w:rsidRPr="008A5523"><w:rPr><w:i/></w:rPr><w:t>Bestronics</w:t></w:r><w:proofErr w:type="spellEnd"/><w:r w:rsidRPr="008A5523"><w:rPr><w:i/></w:rPr><w:t xml:space="preserve"> Inc.</w:t></w:r><w:r><w:tab/></w:r><w:r w:rsidRPr="008A5523"><w:rPr><w:sz w:val="32"/></w:rPr><w:t>MANUFACTURING PROCESSING INSTRUCTIONS</w:t></w:r></w:p></w:hdr>'
			elif (g[a]["_partname"].find("header") > -1):
			#	print "header"
				g[a]["_blob"] = '<?xml version="1.0" encoding="UTF-8" standalone="yes"?><w:hdr xmlns:wpc="http://schemas.microsoft.com/office/word/2010/wordprocessingCanvas" xmlns:mc="http://schemas.openxmlformats.org/markup-compatibility/2006" xmlns:o="urn:schemas-microsoft-com:office:office" xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships" xmlns:m="http://schemas.openxmlformats.org/officeDocument/2006/math" xmlns:v="urn:schemas-microsoft-com:vml" xmlns:wp14="http://schemas.microsoft.com/office/word/2010/wordprocessingDrawing" xmlns:wp="http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing" xmlns:w10="urn:schemas-microsoft-com:office:word" xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" xmlns:w14="http://schemas.microsoft.com/office/word/2010/wordml" xmlns:w15="http://schemas.microsoft.com/office/word/2012/wordml" xmlns:wpg="http://schemas.microsoft.com/office/word/2010/wordprocessingGroup" xmlns:wpi="http://schemas.microsoft.com/office/word/2010/wordprocessingInk" xmlns:wne="http://schemas.microsoft.com/office/word/2006/wordml" xmlns:wps="http://schemas.microsoft.com/office/word/2010/wordprocessingShape" mc:Ignorable="w14 w15 wp14"><w:p w:rsidR="008A5523" w:rsidRDefault="008A5523" w:rsidP="00AF4909"><w:pPr><w:pStyle w:val="Header"/><w:jc w:val="both"/></w:pPr><w:proofErr w:type="spellStart"/><w:r w:rsidRPr="008A5523"><w:rPr><w:i/></w:rPr><w:t>Bestronics</w:t></w:r><w:proofErr w:type="spellEnd"/><w:r w:rsidRPr="008A5523"><w:rPr><w:i/></w:rPr><w:t xml:space="preserve"> Inc.</w:t></w:r><w:r><w:tab/></w:r><w:r w:rsidRPr="008A5523"><w:rPr><w:sz w:val="32"/></w:rPr><w:t>MANUFACTURING PROCESSING INSTRUCTIONS</w:t></w:r></w:p></w:hdr>'
			if (g[a]["_partname"].find("footer1.xml") > -1):
				g[a]["_blob"] = '<?xml version="1.0" encoding="UTF-8" standalone="yes"?><w:ftr xmlns:wpc="http://schemas.microsoft.com/office/word/2010/wordprocessingCanvas" xmlns:mc="http://schemas.openxmlformats.org/markup-compatibility/2006" xmlns:o="urn:schemas-microsoft-com:office:office" xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships" xmlns:m="http://schemas.openxmlformats.org/officeDocument/2006/math" xmlns:v="urn:schemas-microsoft-com:vml" xmlns:wp14="http://schemas.microsoft.com/office/word/2010/wordprocessingDrawing" xmlns:wp="http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing" xmlns:w10="urn:schemas-microsoft-com:office:word" xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" xmlns:w14="http://schemas.microsoft.com/office/word/2010/wordml" xmlns:w15="http://schemas.microsoft.com/office/word/2012/wordml" xmlns:wpg="http://schemas.microsoft.com/office/word/2010/wordprocessingGroup" xmlns:wpi="http://schemas.microsoft.com/office/word/2010/wordprocessingInk" xmlns:wne="http://schemas.microsoft.com/office/word/2006/wordml" xmlns:wps="http://schemas.microsoft.com/office/word/2010/wordprocessingShape" mc:Ignorable="w14 w15 wp14"><w:p w:rsidR="00D6537D" w:rsidRDefault="00D6537D"><w:pPr><w:pStyle w:val="Footer"/></w:pPr><w:r><w:t>' + \
								assyNum + ' Rev ' + assyRev + '</w:t></w:r><w:r><w:tab/></w:r><w:r><w:tab/><w:t xml:space="preserve">Page </w:t></w:r><w:r><w:fldChar w:fldCharType="begin"/></w:r><w:r><w:instrText xml:space="preserve"> PAGE   \* MERGEFORMAT </w:instrText></w:r><w:r><w:fldChar w:fldCharType="separate"/></w:r><w:r><w:rPr><w:noProof/></w:rPr><w:t>1</w:t></w:r><w:r><w:rPr><w:noProof/></w:rPr><w:fldChar w:fldCharType="end"/></w:r><w:r><w:rPr><w:noProof/></w:rPr><w:t xml:space="preserve"> of </w:t></w:r><w:r><w:rPr><w:noProof/></w:rPr><w:fldChar w:fldCharType="begin"/></w:r><w:r><w:rPr><w:noProof/></w:rPr><w:instrText xml:space="preserve"> NUMPAGES  \* Arabic  \* MERGEFORMAT </w:instrText></w:r><w:r><w:rPr><w:noProof/></w:rPr><w:fldChar w:fldCharType="separate"/></w:r><w:r><w:rPr><w:noProof/></w:rPr><w:t>1</w:t></w:r><w:r><w:rPr><w:noProof/></w:rPr><w:fldChar w:fldCharType="end"/></w:r></w:p></w:ftr>'
			elif (g[a]["_partname"].find("footer") > -1):
				g[a]["_blob"] = '<?xml version="1.0" encoding="UTF-8" standalone="yes"?><w:ftr xmlns:wpc="http://schemas.microsoft.com/office/word/2010/wordprocessingCanvas" xmlns:mc="http://schemas.openxmlformats.org/markup-compatibility/2006" xmlns:o="urn:schemas-microsoft-com:office:office" xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships" xmlns:m="http://schemas.openxmlformats.org/officeDocument/2006/math" xmlns:v="urn:schemas-microsoft-com:vml" xmlns:wp14="http://schemas.microsoft.com/office/word/2010/wordprocessingDrawing" xmlns:wp="http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing" xmlns:w10="urn:schemas-microsoft-com:office:word" xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" xmlns:w14="http://schemas.microsoft.com/office/word/2010/wordml" xmlns:w15="http://schemas.microsoft.com/office/word/2012/wordml" xmlns:wpg="http://schemas.microsoft.com/office/word/2010/wordprocessingGroup" xmlns:wpi="http://schemas.microsoft.com/office/word/2010/wordprocessingInk" xmlns:wne="http://schemas.microsoft.com/office/word/2006/wordml" xmlns:wps="http://schemas.microsoft.com/office/word/2010/wordprocessingShape" mc:Ignorable="w14 w15 wp14"><w:p w:rsidR="00D6537D" w:rsidRDefault="00D6537D"><w:pPr><w:pStyle w:val="Footer"/></w:pPr><w:r><w:t>' + \
								assyNum + ' Rev ' + assyRev + '</w:t></w:r><w:r><w:tab/></w:r><w:r><w:tab/><w:t xml:space="preserve">Page </w:t></w:r><w:r><w:fldChar w:fldCharType="begin"/></w:r><w:r><w:instrText xml:space="preserve"> PAGE   \* MERGEFORMAT </w:instrText></w:r><w:r><w:fldChar w:fldCharType="separate"/></w:r><w:r><w:rPr><w:noProof/></w:rPr><w:t>1</w:t></w:r><w:r><w:rPr><w:noProof/></w:rPr><w:fldChar w:fldCharType="end"/></w:r><w:r><w:rPr><w:noProof/></w:rPr><w:t xml:space="preserve"> of </w:t></w:r><w:r><w:rPr><w:noProof/></w:rPr><w:fldChar w:fldCharType="begin"/></w:r><w:r><w:rPr><w:noProof/></w:rPr><w:instrText xml:space="preserve"> NUMPAGES  \* Arabic  \* MERGEFORMAT </w:instrText></w:r><w:r><w:rPr><w:noProof/></w:rPr><w:fldChar w:fldCharType="separate"/></w:r><w:r><w:rPr><w:noProof/></w:rPr><w:t>1</w:t></w:r><w:r><w:rPr><w:noProof/></w:rPr><w:fldChar w:fldCharType="end"/></w:r></w:p></w:ftr>'
		
	doc.save(out_file)